import re
from docx import Document


# pip install python-docx


def markdown_to_word(markdown_text):
    doc = Document()
    rows = markdown_text.strip().split('\n')
    table = doc.add_table(rows=len(rows), cols=rows[0].count('|') - 1)
    table.style = 'Table Grid'

    for i, row in enumerate(rows):
        cells = re.split(r'\s*\|\s*', row.strip())[1:-1]
        for j, cell in enumerate(cells):
            table.cell(i, j).text = cell

    return doc

# Example Markdown table
markdown_table = """
| Header 1 | Header 2 | Header 3 |
|----------|----------|----------|
| Row 1    | Data 1   | Data 2   |
| Row 2    | Data 3   | Data 4   |
"""

doc = markdown_to_word(markdown_table)
doc.save('output.docx')
